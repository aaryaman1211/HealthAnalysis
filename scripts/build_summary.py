"""Rebuild Daily_Health_Summary.xlsx from every export folder in the repo.

Run: python3 scripts/build_summary.py
Then open the workbook in Excel (or run LibreOffice) once so formulas get calculated values.
"""

from datetime import timedelta

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import healthdata as H

OUT = H.ROOT / "Daily_Health_Summary.xlsx"
data = H.build()
daily, sessions, plist = data["daily"], data["sessions"], data["periods"]
names = [p["name"] for p in plist]

F = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F4E78")
HDR_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
MUTED = Font(name=F, size=10, italic=True, color="999999")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILLS = {k: PatternFill("solid", start_color=c) for k, c in
         {"activity": "E2EFDA", "heart": "FCE4D6", "workout": "DDEBF7", "sleep": "E4DFEC", "energy": "FFF2CC"}.items()}
CENTER = Alignment(horizontal="center")


def header(ws, labels, widths):
    for i, (h, w) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32


def clean(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT or v is pd.NA:
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v


def period_label(p):
    return f'{p["name"]}: {p["start"]:%d %b} – {p["end"]:%d %b %Y}'


wb = Workbook()

# ---------------- Daily Summary ----------------
ds = wb.active
ds.title = "Daily Summary"
cols = [
    ("date", "Date", 12, None, "yyyy-mm-dd"),
    ("period", "Period", 8, None, None),
    ("week", "Week Start (Mon)", 12, None, "yyyy-mm-dd"),
    ("steps", "Steps", 9, "activity", "#,##0"),
    ("distance_km", "Distance (km)", 10, "activity", "0.00"),
    ("run_distance_km", "Run Distance (km)", 10, "activity", "0.00"),
    ("act_cal", "Activity Calories (kcal)", 12, "activity", "#,##0"),
    ("active_min", "Active Minutes", 10, "activity", "0"),
    ("rhr_min", "Resting HR (bpm)", 10, "heart", "0"),
    ("avg_hr", "Avg HR (bpm)", 10, "heart", "0.0"),
    ("max_hr", "Max HR (bpm)", 10, "heart", "0"),
    ("sleep_hr_avg", "Sleeping HR (bpm)", 10, "heart", "0.0"),
    ("vigorous_min", f"Vigorous Minutes (HR ≥ {H.VIGOROUS_HR})", 12, "heart", "0"),
    ("hrv", "HRV", 7, "heart", None),
    ("workouts", "Workout(s)", 44, "workout", None),
    ("w_count", "Workout Count", 10, "workout", "0"),
    ("w_min", "Workout Minutes", 10, "workout", "0"),
    ("w_cal", "Workout Calories (kcal)", 12, "workout", "#,##0"),
    ("total_sleep", "Total Sleep (min)", 10, "sleep", "0"),
    ("deep", "Deep Sleep (min)", 10, "sleep", "0"),
    ("light", "Light Sleep (min)", 10, "sleep", "0"),
    ("rem", "REM Sleep (min)", 10, "sleep", "0"),
    ("awake", "Time Awake (min)", 10, "sleep", "0"),
    ("bedtime", "Fell Asleep (IST)", 17, "sleep", "yyyy-mm-dd hh:mm"),
    ("waketime", "Woke Up (IST)", 17, "sleep", "yyyy-mm-dd hh:mm"),
    ("naps", "Naps (count)", 8, "sleep", "0"),
    ("tdee", "Est. Energy Burned (kcal)", 12, "energy", "#,##0"),
]
COL = {k: get_column_letter(i) for i, (k, *_) in enumerate(cols, start=1)}
header(ds, [c[1] for c in cols], [c[2] for c in cols])
for i, row in daily.iterrows():
    r = i + 2
    vals = row.to_dict()
    vals["week"] = f'={COL["date"]}{r}-WEEKDAY({COL["date"]}{r},3)'
    vals["hrv"] = "N/A"
    vals["total_sleep"] = f'=IF(COUNT({COL["deep"]}{r}:{COL["rem"]}{r})=3,{COL["deep"]}{r}+{COL["light"]}{r}+{COL["rem"]}{r},"")'
    for ci, (key, _, _, section, fmt) in enumerate(cols, start=1):
        v = clean(vals.get(key))
        if key == "vigorous_min" or key == "active_min":
            v = int(v) if v is not None else None
        cell = ds.cell(row=r, column=ci, value=v)
        cell.border = BORDER
        cell.font = MUTED if key == "hrv" else (BOLD if key == "date" else BODY)
        if section:
            cell.fill = FILLS[section]
        if key != "workouts":
            cell.alignment = CENTER
        if fmt:
            cell.number_format = fmt
last_row = len(daily) + 1
ds.freeze_panes = "B2"
ds.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last_row}"
comments = {
    "period": "One period per export folder: " + "; ".join(period_label(p) for p in plist) + ".",
    "rhr_min": "Lowest HEARTRATE_AUTO reading of the day. The export has no dedicated resting-HR field.",
    "sleep_hr_avg": "Average heart rate across the night's sleep (SLEEP_MINUTE), attached to the wake-up date.",
    "vigorous_min": f"Minutes with HR ≥ {H.VIGOROUS_HR} bpm (70% of age-predicted max {H.HR_MAX}).",
    "hrv": "Not present in any export file.",
    "bedtime": "Converted from the export's UTC timestamps to IST (UTC+5:30).",
    "tdee": "BMR (Katch-McArdle, 65 kg, 13% body fat) + non-workout step calories + workout calories. See ANALYSIS.md.",
}
for key, text in comments.items():
    ds[f"{COL[key]}1"].comment = Comment(text, "Claude")

rng = lambda key: f"'Daily Summary'!${COL[key]}$2:${COL[key]}${last_row}"

# ---------------- Workout Sessions ----------------
ws = wb.create_sheet("Workout Sessions")
w_hdr = ["Date (IST)", "Period", "Type Code", "Inferred Activity", "Start (IST)", "Duration (min)",
         "Distance (m)", "Avg Pace raw (s/m)", "Pace (min/km)", "Calories (kcal)", "kcal / min", "Avg HR (bpm)", "Max HR (bpm)"]
header(ws, w_hdr, [12, 8, 9, 28, 17, 10, 11, 11, 10, 10, 9, 10, 10])
for i, s in sessions.iterrows():
    r = i + 2
    row = [
        clean(s["date"]), s["period"], int(s["type"]), s["label"], clean(s["start"]),
        round(s["dur_min"], 1), s["dist_m"] if s["dist_m"] > 0 else None,
        s["pace_raw"] if s["pace_raw"] > 0 else None,
        f'=IF(AND(G{r}>={H.MIN_GPS_DISTANCE_M},H{r}>0),H{r}*1000/60,"")',
        s["cal_device"], f'=IF(F{r}>0,J{r}/F{r},"")',
        clean(round(s["avg_hr"], 1)) if pd.notna(s["avg_hr"]) else None,
        clean(s["max_hr"]),
    ]
    for c, v in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font, cell.border = BODY, BORDER
        if c != 4:
            cell.alignment = CENTER
    ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
    ws.cell(row=r, column=5).number_format = "yyyy-mm-dd hh:mm"
    ws.cell(row=r, column=9).number_format = "0.00"
    ws.cell(row=r, column=11).number_format = "0.0"
n_sessions = len(sessions)
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:M{n_sessions + 1}"
ws["I1"].comment = Comment(f"Blank when distance < {H.MIN_GPS_DISTANCE_M} m: those sessions are GPS dropouts, so their pace isn't real.", "Claude")
ws["L1"].comment = Comment("Average of HEARTRATE_AUTO readings inside the session window.", "Claude")
wrng = lambda c: f"'Workout Sessions'!${c}$2:${c}${n_sessions + 1}"

# ---------------- Period Comparison ----------------
pc = wb.create_sheet("Period Comparison")
pc_cols = ["Metric"] + [period_label(p) for p in plist]
if len(plist) >= 2:
    pc_cols += [f"Change ({names[-2]} → {names[-1]})", "% Change"]
header(pc, pc_cols, [34] + [22] * len(plist) + [16, 11][: len(pc_cols) - 1 - len(plist)])
P = rng("period")
avg = lambda key, scale="": (lambda p: f'=AVERAGEIFS({rng(key)},{P},"{p}"){scale}')
metrics = [
    ("Days", lambda p: f'=COUNTIF({P},"{p}")', "0"),
    ("Avg steps / day", avg("steps"), "#,##0"),
    ("Avg distance / day (km)", avg("distance_km"), "0.00"),
    ("Avg active minutes / day", avg("active_min"), "0"),
    ("Avg activity calories / day", avg("act_cal"), "#,##0"),
    ("Avg resting HR (bpm)", avg("rhr_min"), "0.0"),
    ("Avg sleeping HR (bpm)", avg("sleep_hr_avg"), "0.0"),
    ("Avg daily HR (bpm)", avg("avg_hr"), "0.0"),
    ("Avg vigorous minutes / day", avg("vigorous_min"), "0.0"),
    ("Workout days", lambda p: f'=COUNTIFS({P},"{p}",{rng("w_count")},">0")', "0"),
    ("Workout sessions", lambda p: f'=SUMIFS({rng("w_count")},{P},"{p}")', "0"),
    ("Avg workout minutes / day", avg("w_min"), "0"),
    ("Avg workout calories / day", avg("w_cal"), "#,##0"),
    ("Avg sleep (hours)", avg("total_sleep", "/60"), "0.00"),
    ("Nights under 6 h", lambda p: f'=COUNTIFS({P},"{p}",{rng("total_sleep")},"<360")', "0"),
    ("Avg deep sleep (min)", avg("deep"), "0"),
    ("Avg REM sleep (min)", avg("rem"), "0"),
    ("Deep sleep % of total", lambda p: f'=SUMIFS({rng("deep")},{P},"{p}")/SUMIFS({rng("total_sleep")},{P},"{p}")', "0.0%"),
    ("Runs logged", lambda p: f'=COUNTIFS({wrng("B")},"{p}",{wrng("C")},1)', "0"),
    ("Total run distance (km)", lambda p: f'=SUMIFS({wrng("G")},{wrng("B")},"{p}",{wrng("C")},1)/1000', "0.0"),
    ("Avg run pace (min/km)", lambda p: f'=AVERAGEIFS({wrng("I")},{wrng("B")},"{p}",{wrng("C")},1)', "0.00"),
    ("Avg est. energy burned / day (kcal)", avg("tdee"), "#,##0"),
]
n = len(plist)
for i, (name, f, fmt) in enumerate(metrics, start=2):
    pc.cell(row=i, column=1, value=name).font = BOLD
    for j, p in enumerate(names):
        c = pc.cell(row=i, column=2 + j, value=f(p))
        c.number_format, c.font, c.alignment = fmt, BODY, CENTER
    if n >= 2:
        a, b = get_column_letter(n), get_column_letter(n + 1)
        ch = pc.cell(row=i, column=n + 2, value=f"={b}{i}-{a}{i}")
        ch.number_format = "0.0%" if fmt == "0.0%" else "+#,##0.0;-#,##0.0;0"
        pct = pc.cell(row=i, column=n + 3, value=f'=IF({a}{i}=0,"",({b}{i}-{a}{i})/{a}{i})')
        pct.number_format = "+0.0%;-0.0%;0.0%"
    for c in range(1, len(pc_cols) + 1):
        pc.cell(row=i, column=c).border = BORDER
        if c > 1 + n:
            pc.cell(row=i, column=c).font, pc.cell(row=i, column=c).alignment = BODY, CENTER
pc.cell(row=len(metrics) + 3, column=1,
        value="Lower is better for resting/sleeping HR and run pace (min/km). All values are formulas over the other tabs. See report/REPORT.md for significance tests.").font = MUTED
pc.freeze_panes = "B2"

# ---------------- Weekly Trends ----------------
wt = wb.create_sheet("Weekly Trends")
header(wt, ["Week Start (Mon)", "Days", "Avg Steps", "Avg Resting HR", "Avg Sleeping HR", "Avg Sleep (h)",
            "Workout Sessions", "Avg Workout kcal / day", "Vigorous Minutes"], [14, 7, 11, 11, 11, 11, 11, 13, 11])
first, last = daily["date"].min(), daily["date"].max()
weeks, wk = [], first - timedelta(days=first.weekday())
while wk <= last:
    if sum(first <= wk + timedelta(days=k) <= last for k in range(7)) >= 4:
        weeks.append(wk)
    wk += timedelta(days=7)
W = rng("week")
for r, wk in enumerate(weeks, start=2):
    row = [
        wk.to_pydatetime(), f"=COUNTIF({W},A{r})",
        f"=AVERAGEIFS({rng('steps')},{W},A{r})",
        f"=AVERAGEIFS({rng('rhr_min')},{W},A{r})",
        f"=AVERAGEIFS({rng('sleep_hr_avg')},{W},A{r})",
        f"=AVERAGEIFS({rng('total_sleep')},{W},A{r})/60",
        f"=SUMIFS({rng('w_count')},{W},A{r})",
        f"=AVERAGEIFS({rng('w_cal')},{W},A{r})",
        f"=SUMIFS({rng('vigorous_min')},{W},A{r})",
    ]
    for c, v in enumerate(row, start=1):
        cell = wt.cell(row=r, column=c, value=v)
        cell.font, cell.border, cell.alignment = BODY, BORDER, CENTER
    for c, fmt in {1: "yyyy-mm-dd", 3: "#,##0", 4: "0.0", 5: "0.0", 6: "0.00", 8: "#,##0", 9: "0"}.items():
        wt.cell(row=r, column=c).number_format = fmt
wt_last = len(weeks) + 1
wt.cell(row=wt_last + 2, column=1, value="Weeks with fewer than 4 days of data are left out so partial weeks don't distort the trend.").font = MUTED
wt.freeze_panes = "A2"


def add_chart(title, col, y_title, anchor):
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.x_axis.title = title, y_title, "Week starting"
    ch.height, ch.width, ch.legend = 7.5, 17, None
    ch.add_data(Reference(wt, min_col=col, min_row=1, max_row=wt_last), titles_from_data=True)
    ch.set_categories(Reference(wt, min_col=1, min_row=2, max_row=wt_last))
    ch.x_axis.number_format = "dd mmm"
    ch.x_axis.delete = ch.y_axis.delete = False
    wt.add_chart(ch, anchor)


add_chart("Avg daily steps by week", 3, "Steps", "K2")
add_chart("Avg resting HR by week", 4, "bpm", "K18")
add_chart("Avg sleep by week", 6, "Hours", "K34")

# ---------------- Notes ----------------
nt = wb.create_sheet("Notes")
nt.column_dimensions["A"].width = 115
weights = "; ".join(f'{t:%Y-%m-%d}: {w} kg' for t, w in zip(data["body"]["time"], data["body"]["weight"]))
notes = [
    ("Data sources & assumptions", "title"),
    (f"Covers {first:%Y-%m-%d} to {last:%Y-%m-%d} ({len(daily)} days), merged from every 'Health Data*' export folder. "
     + "Periods: " + "; ".join(period_label(p) for p in plist) + ". A date that appears in two exports belongs to the earlier one; duplicate rows are de-duplicated with the newer export winning.", None),
    ("Timezone: all times are IST (UTC+5:30). HEARTRATE_AUTO, ACTIVITY_MINUTE, ACTIVITY_STAGE and SLEEP_MINUTE are already local; SPORT start times and SLEEP start/stop are UTC in the export and are converted. The offset was confirmed by aligning workouts with heart-rate spikes (79 of 129 sessions of 20+ minutes peak at exactly +5:30).", None),
    ("Workouts are assigned to their IST calendar date.", None),
    ("Resting HR = lowest HEARTRATE_AUTO reading of the day. Sleeping HR = average HR across that night's sleep from SLEEP_MINUTE. There is no dedicated resting-HR field.", None),
    (f"Vigorous minutes = minutes with HR ≥ {H.VIGOROUS_HR} bpm (70% of an age-predicted max of {H.HR_MAX} for age {H.AGE}).", None),
    ("HRV and respiratory rate are not recorded (the respiratory_rate column is empty in every export).", None),
    ("Active Minutes = total duration of ACTIVITY_STAGE movement segments for that date.", None),
    ("Activity Calories (device daily figure) already includes workout movement, so it overlaps Workout Calories; don't add them together.", None),
    ("Workout names are inferred from Zepp's numeric type codes by distance, pace and intensity; the export has no names.", None),
    (f"Run pace (min/km) = SPORT avgPace (seconds per metre) × 1000 / 60, blank for sessions under {H.MIN_GPS_DISTANCE_M} m (GPS dropouts, e.g. 2026-06-25 run: 59 min for 58 m).", None),
    ("Sleep fields come from SLEEP; Total Sleep = Deep + Light + REM. The sleep row's date is the wake-up date.", None),
    ("Naps: every nap timestamp in the exports is a placeholder (2024-10-04), so only the count is shown.", None),
    (f"Est. Energy Burned = BMR {data['bmr']:.0f} kcal (Katch-McArdle at {H.WEIGHT_KG:.0f} kg, {H.BODY_FAT:.0%} body fat, user-provided) + non-workout step calories + workout calories, with device calories rescaled from the export's profile weight to {H.WEIGHT_KG:.0f} kg.", None),
    (f"Body weight readings on record: {weights}. The 65 kg entry may be a profile update rather than a scale reading.", None),
    ("Rebuild: python3 scripts/build_summary.py, then open the file in Excel once so formulas get calculated values.", None),
]
for i, (line, kind) in enumerate(notes, start=1):
    c = nt.cell(row=i, column=1, value=line)
    c.font = Font(name=F, size=13, bold=True) if kind == "title" else BODY
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print(f"saved {OUT}: {len(daily)} days, {n_sessions} sessions, {len(weeks)} weeks, periods {names}")
